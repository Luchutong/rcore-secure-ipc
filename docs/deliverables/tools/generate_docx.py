#!/usr/bin/env python3
"""Convert COURSE_REPORT.md into an editable, conservative Word draft."""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "COURSE_REPORT.md"
OUTPUT = ROOT / "generated" / "rcore-secure-ipc-course-report-draft.docx"
FONT_CN = "SimSun"
FONT_HEADING = "Microsoft YaHei"


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def set_run_font(run, chinese=FONT_CN, latin="Times New Roman", size=None, bold=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), chinese)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, rows):
    if len(rows) < 2:
        return
    data = [row for index, row in enumerate(rows) if index != 1]
    width = max(len(row) for row in data)
    table = doc.add_table(rows=len(data), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(data):
        for j in range(width):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = clean_inline(row[j]) if j < len(row) else ""
            cell.text = text
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=9.5, bold=(i == 0))
            if i == 0:
                shade_cell(cell, "E8EDF3")
    doc.add_paragraph()


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.3)
    set_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_after = Pt(0)

    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = FONT_HEADING
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEADING)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(20, 20, 20)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("课程设计报告")
    set_run_font(r, FONT_HEADING, size=24, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    r = p.add_run("基于 rCore 的安全进程间通信机制设计与实现")
    set_run_font(r, FONT_HEADING, size=20, bold=True)
    for _ in range(4):
        doc.add_paragraph()
    fields = [
        "课程名称：    [待填写]",
        "学院专业：    [待填写]",
        "班    级：    [待填写]",
        "组    员：    [姓名与学号待填写]",
        "指导教师：    [待填写]",
        "完成日期：    2026 年 9 月",
    ]
    for field in fields:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(field)
        set_run_font(r, FONT_CN, size=14)
    doc.add_page_break()


def parse_report(doc, lines):
    i = 0
    paragraph_buffer = []
    in_code = False
    code_lines = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = clean_inline("".join(part.strip() for part in paragraph_buffer))
            if text:
                doc.add_paragraph(text)
            paragraph_buffer = []

    while i < len(lines):
        raw = lines[i].rstrip("\n")
        stripped = raw.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.right_indent = Cm(0.7)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.0
                r = p.add_run("\n".join(code_lines))
                set_run_font(r, "DejaVu Sans Mono", "Courier New", 9)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped == "---":
            flush_paragraph()
            i += 1
            continue

        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            flush_paragraph()
            image_path = ROOT / image_match.group(2)
            if image_path.suffix.lower() == ".svg":
                image_path = image_path.with_suffix(".png")
            if image_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.add_run().add_picture(str(image_path), width=Cm(15.2))
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = clean_inline(heading.group(2))
            if level == 1:
                # The real title is already on the cover.
                i += 1
                continue
            doc.add_heading(text, level=min(level - 1, 3))
            i += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            p = doc.add_paragraph(clean_inline(stripped.lstrip("> ")))
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(90, 90, 90)
            i += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or numbered:
            flush_paragraph()
            text = clean_inline((bullet or numbered).group(1))
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            p.paragraph_format.first_line_indent = Cm(0)
            p.add_run(text)
            i += 1
            continue

        paragraph_buffer.append(raw + " ")
        i += 1

    flush_paragraph()


def build():
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "基于 rCore 的安全进程间通信机制设计与实现"
    doc.core_properties.subject = "课程设计报告可编辑草稿"
    doc.core_properties.comments = "Generated from verified v1.0.0 delivery materials"
    add_cover(doc)
    parse_report(doc, SOURCE.read_text(encoding="utf-8").splitlines())
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
